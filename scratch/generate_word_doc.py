import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def create_styled_document():
    doc = Document()

    # Page setup - 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    md_path = r"f:\Student_OS\PROJECT_DOCUMENTATION.md"
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    i = 0
    in_code_block = False
    code_lines = []
    code_lang = ""

    while i < len(lines):
        line = lines[i]

        # Fenced code block handling
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lang = line.strip().replace('```', '')
                code_lines = []
            else:
                in_code_block = False
                # Render code block table
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False
                table.columns[0].width = Inches(7.0)
                cell = table.cell(0, 0)
                set_cell_background(cell, "F8F9FA")
                set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
                
                cell_p = cell.paragraphs[0]
                cell_p.paragraph_format.space_before = Pt(2)
                cell_p.paragraph_format.space_after = Pt(2)
                cell_p.paragraph_format.line_spacing = 1.15
                
                full_code = '\n'.join(code_lines)
                run = cell_p.add_run(full_code)
                run.font.name = 'Consolas'
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0x2B, 0x30, 0x3A)
                
                # Add spacing after table
                p_after = doc.add_paragraph()
                p_after.paragraph_format.space_before = Pt(2)
                p_after.paragraph_format.space_after = Pt(4)

            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Header 1
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[2:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0A, 0x36, 0x63)
            i += 1
            continue

        # Header 2
        if line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[3:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1B, 0x4D, 0x89)
            i += 1
            continue

        # Header 3
        if line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[4:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x28, 0x5C, 0x9B)
            i += 1
            continue

        # Header 4
        if line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[5:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            i += 1
            continue

        # Blockquote
        if line.startswith('> '):
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.columns[0].width = Inches(7.0)
            cell = table.cell(0, 0)
            set_cell_background(cell, "F0F4F8")
            set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
            
            # Left border thick accent
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="0A3663"/>'
                f'  <w:top w:val="none"/>'
                f'  <w:bottom w:val="none"/>'
                f'  <w:right w:val="none"/>'
                f'</w:tcBorders>'
            )
            tcPr.append(tcBorders)

            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.space_after = Pt(2)
            # Remove >
            bq_text = line[2:].strip()
            # Parse bold/italic
            add_formatted_text(cp, bq_text, italic_all=True)
            
            p_after = doc.add_paragraph()
            p_after.paragraph_format.space_before = Pt(2)
            p_after.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # Markdown Table Detection
        if line.strip().startswith('|') and '|' in line[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1

            if len(table_lines) >= 2:
                render_markdown_table(doc, table_lines)
            continue

        # Horizontal Rule
        if line.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            # Subtle horizontal rule
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="D3D9DE"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue

        # Bullet list
        if line.strip().startswith('- ') or line.strip().startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            indent = len(line) - len(line.lstrip())
            if indent >= 4:
                p.paragraph_format.left_indent = Inches(0.4)
            elif indent >= 2:
                p.paragraph_format.left_indent = Inches(0.25)
            bullet_text = line.strip()[2:]
            add_formatted_text(p, bullet_text)
            i += 1
            continue

        # Numbered list
        num_match = re.match(r'^(\d+)\.\s+(.*)', line.strip())
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, num_match.group(2))
            i += 1
            continue

        # Regular Paragraph
        if line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, line.strip())
        
        i += 1

    output_path = r"f:\Student_OS\PROJECT_DOCUMENTATION.docx"
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

def add_formatted_text(paragraph, text, italic_all=False):
    # Regex for bold: **text**, code: `code`, links: [title](url), italic: *text*
    pattern = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*]+\*)')
    parts = pattern.split(text)

    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            if italic_all:
                run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B)
            if italic_all:
                run.italic = True
        elif part.startswith('[') and '](' in part and part.endswith(')'):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', part)
            if m:
                label, url = m.groups()
                run = paragraph.add_run(f"{label} ({url})")
                run.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
                run.underline = True
                if italic_all:
                    run.italic = True
            else:
                run = paragraph.add_run(part)
                if italic_all:
                    run.italic = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
            if italic_all:
                run.italic = True

def render_markdown_table(doc, lines):
    # Filter header separator line (|:---|:---|)
    parsed_rows = []
    for line in lines:
        raw_cells = [c.strip() for c in line.strip('|').split('|')]
        # Check if separator row
        if all(re.match(r'^:?-+:?$', c) for c in raw_cells):
            continue
        parsed_rows.append(raw_cells)

    if not parsed_rows:
        return

    col_count = max(len(r) for r in parsed_rows)
    table = doc.add_table(rows=len(parsed_rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    total_width = Inches(7.0)
    col_width = total_width / col_count

    for r_idx, row in enumerate(parsed_rows):
        # Set cantSplit on row
        trPr = table.rows[r_idx]._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        # Set header repeat if header row
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for c_idx in range(col_count):
            cell = table.cell(r_idx, c_idx)
            cell.width = col_width
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

            if r_idx == 0:
                set_cell_background(cell, "0A3663")
            elif r_idx % 2 == 1:
                set_cell_background(cell, "FFFFFF")
            else:
                set_cell_background(cell, "F8FAFC")

            cell_p = cell.paragraphs[0]
            cell_p.paragraph_format.space_before = Pt(1)
            cell_p.paragraph_format.space_after = Pt(1)
            cell_p.paragraph_format.line_spacing = 1.1

            text = row[c_idx] if c_idx < len(row) else ""
            if r_idx == 0:
                run = cell_p.add_run(text)
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(9.5)
            else:
                add_formatted_text(cell_p, text)
                for r in cell_p.runs:
                    if r.font.size is None:
                        r.font.size = Pt(9.0)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(4)

if __name__ == '__main__':
    create_styled_document()
